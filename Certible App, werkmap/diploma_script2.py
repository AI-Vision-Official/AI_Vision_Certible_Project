import os
import re
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# === Aanpassen naar jouw map ===
root_dir = r"L:\3. Study\@Diploma's Definitief"  # <-- Pas dit aan indien nodig
output_file = "L:/GIT/Projects/Certible Project/certible_export.xlsx"

# Rubriek bepalen op basis van submapnaam
def detect_rubric(folder):
    name = os.path.basename(folder)
    if " - " in name:
        return name.split(" - ")[0].replace("Certificates ", "").strip()
    return "Other"

# Herkenbare patronen
levels = ["MBO", "HBO", "WO", "Bachelor", "Master", "PhD"]
languages = ["English", "Dutch", "German", "French", "Spanish"]
schools = ["NTI", "LOI", "OpenUniversity", "ROC", "Cambridge", "Coursera", "Udemy", "FEMA", "University", "Iowa"]

# Info uit bestandsnaam halen
def detect_info(filename):
    school = next((s for s in schools if s.lower() in filename.lower()), "")
    level = next((l for l in levels if l.lower() in filename.lower()), "")
    language = next((lang for lang in languages if lang.lower() in filename.lower()), "")
    match = re.search(r'(20\d{2})', filename)
    valid_until = f"{match.group(1)}-12-31" if match else ""
    return school, level, language, valid_until

# Werkblad instellen
def setup_worksheet(ws):
    headers = [
        "Certificate_name", "School-Issuer", "Subject", "Part of Series", "Diploma Number", "Date credential",
        "Level", "Language", "Location", "Certification Type", "Accredited?", "Accreditiation Institute", "Valid Until", "Renewed?", "Status", "Renewed On", "Badge", "Badge Platform" "Remarks"
    ]
    fill = PatternFill(start_color="B7D7F7", end_color="B7D7F7", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center")
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = fill
        cell.font = bold
        cell.alignment = center
        ws.column_dimensions[cell.column_letter].width = 20
    dv = DataValidation(type="list", formula1='"Yes,No"', showDropDown=True)
    ws.add_data_validation(dv)
    dv.add("J2:J10000")

# Nieuw werkboek maken
wb = Workbook()
wb.remove(wb.active)
rubric_counter = {}

for foldername, subfolders, filenames in os.walk(root_dir):
    rel_depth = foldername[len(root_dir):].count(os.sep)
    if rel_depth > 2:
        continue
    rubric = detect_rubric(foldername)
    if rubric not in wb.sheetnames:
        ws = wb.create_sheet(rubric)
        setup_worksheet(ws)
    else:
        ws = wb[rubric]

    for f in filenames:
        if f.lower().endswith((".pdf", ".jpg", ".jpeg", ".png")):
            full_path = os.path.join(foldername, f)
            rel_folder = os.path.relpath(foldername, root_dir)
            school, level, language, valid_until = detect_info(f)
            ws.append([
                f,
                rel_folder,
                full_path,
                "", language, school, level, valid_until,
                "", "", ""
            ])
            rubric_counter[rubric] = rubric_counter.get(rubric, 0) + 1

# Dashboard toevoegen
dashboard = wb.create_sheet("Dashboard")
dashboard.append(["Rubric", "Count"])
for rubric, count in sorted(rubric_counter.items()):
    dashboard.append([rubric, count])
dashboard.column_dimensions["A"].width = 30
dashboard.column_dimensions["B"].width = 10
dashboard["A1"].font = dashboard["B1"].font = Font(bold=True)

wb.save(output_file)
import pandas as pd

# Zet Excel om naar CSV
excel_file = r"L:/GIT/Projects/Certible Project/certible_export.xlsx"
csv_file = r"L:/GIT/Projects/Certible Project/certible_export.csv"

# Lees Excel en schrijf als CSV
df = pd.read_excel(excel_file)
df.to_csv(csv_file, index=False, encoding='utf-8-sig')

import csv
import json

# Pad naar jouw CSV-bestand
input_csv = "L:/GIT/Projects/Certible Project/certible_export.csv"
output_json = "L:/GIT/Projects/Certible Project/certible_export.json"

# Lezen en omzetten
with open(input_csv, mode='r', encoding='latin1') as csv_file:
    reader = csv.DictReader(csv_file)
    data = [row for row in reader]

with open(output_json, mode='w', encoding='utf-8') as json_file:
    json.dump(data, json_file, indent=4)
print(f"✅ Done! File saved to: {output_file}")
print("✅ Excel succesvol omgezet naar CSV!")
print("✅ CSV succesvol omgezet naar JSON!")



