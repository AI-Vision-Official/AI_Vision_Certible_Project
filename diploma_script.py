import sys
print(sys.executable)
import os
import re
import json
import csv
import hashlib
import openpyxl
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# === INSTELLINGEN ===
root_dir = r"L:\3. Study\@Diploma's Definitief"
export_base = r"L:/GIT/Projects/Certible Project"
output_xlsx = os.path.join(export_base, "certible_export.xlsx")
output_csv = os.path.join(export_base, "certible_export.csv")
output_json = os.path.join(export_base, "certible_export.json")
hash_file = os.path.join(export_base, "certible_hashes.json")

# === HULPFUNCTIES ===
def detect_rubric(folder):
    name = os.path.basename(folder)
    if " - " in name:
        return name.split(" - ")[0].replace("Certificates ", "").strip()
    return "Other"

def detect_info(filename):
    levels = ["MBO", "HBO", "WO", "Bachelor", "Master", "PhD"]
    languages = ["English", "Dutch", "German", "French", "Spanish"]
    schools = ["NTI", "LOI", "OpenUniversity", "ROC", "Cambridge", "Coursera", "Udemy", "FEMA", "University", "Iowa"]
    school = next((s for s in schools if s.lower() in filename.lower()), "")
    level = next((l for l in levels if l.lower() in filename.lower()), "")
    language = next((lang for lang in languages if lang.lower() in filename.lower()), "")
    match = re.search(r'(20\d{2})', filename)
    valid_until = f"{match.group(1)}-12-31" if match else ""
    return school, level, language, valid_until

def safe_sheet_title(name):
    return name[:31] if len(name) > 31 else name

def file_hash(path):
    with open(path, 'rb') as f:
        return hashlib.md5(f.read()).hexdigest()

def setup_worksheet(ws):
    headers = [
        "Certificate_name", "School-Issuer", "Subject", "Part of Series", "Diploma Number", "Date credential", "Level", "Language", "Location", "Certification Type", "Accredited?", "Accreditiation Institute", "Valid Until", "Renewed?", "Status", "Renewed On", "Badge", "Badge Platform", "Remarks", "Rubric"
    ]
    fill = PatternFill(start_color="B7D7F7", end_color="B7D7F7", fill_type="solid")
    bold = Font(bold=True)
    center = Alignment(horizontal="center")
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = fill
        cell.font = bold
        cell.alignment = center
        ws.column_dimensions[cell.column_letter].width = 20
    dv = DataValidation(type="list", formula1='"Yes,No"', showDropDown=True)
    ws.add_data_validation(dv)
    dv.add("J2:J10000")

# === INLADEN OUDE HASHES ===
try:
    with open(hash_file, 'r') as f:
        old_hashes = json.load(f)
except FileNotFoundError:
    old_hashes = {}

new_hashes = {}
data_rows = {}
rubric_counter = {}

# === MAAK NIEUW WERKBOEK ===
wb = Workbook()
wb.remove(wb.active)
rubric_counter = {}

for foldername, subfolders, filenames in os.walk(root_dir):
    if foldername[len(root_dir):].count(os.sep) > 2:
        continue

    rubric = detect_rubric(foldername)
    safe_rubric = safe_sheet_title(rubric)
    if safe_rubric not in wb.sheetnames:
        ws = wb.create_sheet(safe_rubric)
        setup_worksheet(ws)
        data_rows[safe_rubric] = 1  # start op rij 2
    else:
        ws = wb[safe_rubric]

    for f in filenames:
        if not f.lower().endswith((".pdf", ".jpg", ".jpeg", ".png")):
            continue

        full_path = os.path.join(foldername, f)
        rel_folder = os.path.relpath(foldername, root_dir)

        school, level, language, valid_until = detect_info(f)
        ws.append([
            f, school, "", "", "", "", level, language, rel_folder, "", "", "", valid_until, "", "", "", "", "", ""
        ])
        rubric_counter[rubric] = rubric_counter.get(rubric, 0) + 1
        data_rows[safe_rubric] += 1

# === DASHBOARD ===
dash = wb.create_sheet("Dashboard")
dash.append(["Rubric", "Count"])
for rubric, count in sorted(rubric_counter.items()):
    dash.append([rubric, count])
dash["A1"].font = dash["B1"].font = Font(bold=True)
dash.column_dimensions["A"].width = 30
dash.column_dimensions["B"].width = 10

# === OPSLAAN XLSX ===
wb.save(output_xlsx)
print(f"✅ Done! File saved to: {output_xlsx}")

# === EXCEL NAAR CSV ===
df = pd.read_excel(output_xlsx)
df.to_csv(output_csv, index=False, encoding='utf-8-sig')
print("✅ Excel succesvol omgezet naar CSV!")

# === CSV NAAR JSON ===
with open(output_csv, mode='r', encoding='latin1') as csv_file:
    reader = csv.DictReader(csv_file)
    data = [row for row in reader]

with open(output_json, mode='w', encoding='utf-8') as json_file:
    json.dump(data, json_file, indent=4)
print("✅ CSV succesvol omgezet naar JSON!")

excel_file = r"L:/GIT/Projects/Certible Project/certible_export.xlsx"
output_json = r"L:/GIT/Projects/Certible Project/certible_export.json"

wb = openpyxl.load_workbook(excel_file)
full_data = []

for sheet in wb.sheetnames:
    if sheet.lower() == "dashboard":
        continue  # skip dashboard

    ws = wb[sheet]
    headers = [cell.value for cell in ws[1]]

    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(cell is None for cell in row):
            continue  # skip empty rows
        row_dict = dict(zip(headers, row))
        row_dict["Rubric"] = sheet  # voeg het tabblad toe als veld
        full_data.append(row_dict)

# Schrijf als JSON
with open(output_json, "w", encoding="utf-8") as f:
    json.dump(full_data, f, indent=4)

print("✅ JSON met rubriekinfo succesvol aangemaakt.")

# === NIEUWE HASHES OPSLAAN ===
with open(hash_file, 'w') as f:
    json.dump(new_hashes, f, indent=2)
